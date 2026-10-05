from pathlib import Path
import json,csv,fnmatch
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
out=Path('deliverables/sonarqube_fix')
rows=json.loads((out/'findings.json').read_text())[1:]
patterns=next(l.split('=',1)[1].split(',') for l in Path('sonar-project.properties').read_text().splitlines() if l.startswith('sonar.exclusions='))
def excluded(p):return any(fnmatch.fnmatch(p,x) or (x.startswith('**/') and fnmatch.fnmatch(p,x[3:])) for x in patterns)
related={'Web:S7926','Web:S6848'}
def status(r):
 if excluded(r[3]): return 'Excluded by pre-existing scan configuration; not repaired'
 if r[2] in ('BUG','VULNERABILITY'): return 'Source fix implemented; SonarQube rescan pending'
 if r[3]=='templates/dashboard.html' and (r[5] in related or (r[5]=='Web:S6819' and r[4] in ('184','187'))): return 'Related source fix implemented; SonarQube rescan pending'
 return 'Code smell not addressed in this bugfix pass'
with (out/'finding_disposition.csv').open('w',newline='',encoding='utf-8-sig') as f:
 w=csv.writer(f);w.writerow(['Workbook row','Severity','Type','File','Original line','Rule','Message','Disposition'])
 for i,r in enumerate(rows,2): w.writerow([i,r[1],r[2],r[3],r[4],r[5],r[6],status(r)])
d=Document();sec=d.sections[0];sec.top_margin=sec.bottom_margin=Inches(.65);sec.left_margin=sec.right_margin=Inches(.75)
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2']:
 st=d.styles[name];st.font.name='Calibri';st.font.color.rgb=RGBColor(0,0,0)
d.styles['Normal'].font.size=Pt(10.5);d.styles['Normal'].paragraph_format.space_after=Pt(5)
d.styles['Title'].font.size=Pt(25)
d.styles['Heading 1'].font.size=Pt(16);d.styles['Heading 2'].font.size=Pt(12)
def p(t):d.add_paragraph(t)
def h(t):d.add_heading(t,1)
def table(headers,records,widths=None):
 t=d.add_table(rows=1,cols=len(headers));t.autofit=False
 if widths:
  for col,width in zip(t.columns,widths): col.width=Inches(width)
 for c,txt in zip(t.rows[0].cells,headers):c.text=txt
 for row in records:
  for c,txt in zip(t.add_row().cells,row):c.text=str(txt)
 for i,row in enumerate(t.rows):
  trpr=row._tr.get_or_add_trPr();trpr.append(OxmlElement('w:cantSplit'))
  if i==0:trpr.append(OxmlElement('w:tblHeader'))
  for j,c in enumerate(row.cells):
   if widths:c.width=Inches(widths[j])
   pr=c._tc.get_or_add_tcPr();b=OxmlElement('w:tcBorders')
   for edge in ['top','left','bottom','right']:
    e=OxmlElement('w:'+edge);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');b.append(e)
   pr.append(b)
   if i==0:
    sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'DCE6F1');pr.append(sh)
   for par in c.paragraphs:
    par.paragraph_format.space_after=Pt(2);par.paragraph_format.space_before=Pt(2)
    for run in par.runs:run.font.size=Pt(9);run.bold=i==0
 return t

d.add_heading('MasterHub IoT Bugfix Report',0)
p('Noushith project  |  5 October 2026')
p('Source workbook: MasterHub_IoT_Noushith_SonarQube_Complete_Analysis.xlsx')
p('Source changes address all 16 BUG findings and all 3 VULNERABILITY findings in the current SonarQube scan scope. Validation completed with 215 passing Python tests and 14 passing JavaScript tests. A fresh SonarQube scan is required before marking these findings closed.')
h('Scope and reconciliation')
p('The supplied workbook contains 4,805 findings across 399 files. Its scan included browser runtime files and historical backups. The existing sonar-project.properties already excludes those directories and generated artifacts. These exclusions were preserved; excluded findings were not repaired or counted as fixed.')
table(['Finding type','Workbook','Current scope','Already excluded'],[['Bugs',304,16,288],['Vulnerabilities',6,3,3],['Code smells',4495,621,3874],['Total',4805,640,4165]],[2.2,1.25,1.35,1.6])
h('Outcome')
p('Implemented source corrections for 19 bug and vulnerability findings, plus 6 related dashboard code smells. The remaining 615 in-scope code smells were not addressed in this pass. The 288 excluded bugs comprise 286 browser/runtime findings and 2 historical-backup findings.')
h('Deployment change')
p('Running python app.py now binds to 127.0.0.1 by default. For mobile, MQTT-associated HTTP clients, or other remote clients that access this Flask server, set MASTERHUB_HOST to the PC\'s LAN interface address before startup. The Flask CLI and external WSGI servers retain their own bind configuration. Debug mode remains disabled.')
p('The two dashboards now use their existing local/system font fallbacks instead of loading Google Fonts stylesheets. Typography may differ on machines without Inter or JetBrains Mono installed.')
d.add_page_break()
h('Dashboard fixes')
d.add_heading('Unique status identifiers',2)
p('Removed the hidden duplicate connection matrix from dashboard.html. The six affected identifiers now occur only once, preventing ambiguous DOM lookups and duplicate accessibility content. The visible sidebar retains its original IDs and event handlers.')
d.add_heading('Keyboard controls',2)
p('Converted the three dismiss backdrops, three connection rows, BCI status pill, and theme toggle into native buttons. Enter and Space use native activation. Removed the theme toggle\'s custom keyboard handler to prevent double activation. Added visible focus outlines and preserved full-width connection rows and full-screen backdrop geometry. Closed backdrops are hidden from keyboard focus.')
d.add_heading('Live sidebar status',2)
p('The retained sidebar labels previously remained at CHECKING. MQTT, USB, BCI, system health, and FSM mode now update from the existing polling requests. Failed requests replace stale values with Unavailable. This additional functional correction is covered by a JavaScript regression test and is not counted as an extra workbook finding.')
d.add_heading('Related accessibility cleanup',2)
p('Removed viewport restrictions that disabled zoom. Native controls also address three non-native-interaction findings and two button-role findings. Existing close controls, Escape handling, IDs, and click actions are preserved.')
h('Backend and diagram fixes')
d.add_heading('Explicit server interface',2)
p('app.py reads MASTERHUB_HOST with a localhost fallback, including blank or whitespace-only values. .env.example documents the setting and the startup log reports the selected interface. This removes the unconditional all-interface default flagged by python:S8392.')
d.add_heading('Device selection',2)
p('core/devices.py now assigns the normalized target once while holding the existing lock. Registered targets and explicit targets that have not yet been discovered remain valid. This removes identical branches without changing routing semantics.')
d.add_heading('Diagram arrow rendering',2)
p('The robotics diagram helper always supplies arc3 with the requested radius. Radius zero already represents a straight connector, so the float equality comparison and alternate branch are unnecessary.')
d.add_heading('External stylesheets',2)
p('Removed the remote Google Fonts stylesheet from dashboard.html and cortex_classic.html. Both pages retain their existing font-family fallback stacks, eliminating the two reported unverified external stylesheet loads.')
d.add_page_break()
h('Validation and remaining work')
table(['Validation','Result'],[['Full Python regression suite','215 passed in 6.88 seconds'],['JavaScript regression suites','14 passed'],['New Python regression cases','10 cases included in the full suite'],['New JavaScript regression case','Success and failure of sidebar status refresh'],['SonarQube rescan','Not run; scanner command unavailable'],['Physical device and browser checks','Live hardware and interactive browser checks not performed']],[2.7,3.7])
p('Python tests cover rendered dashboard ID uniqueness and native controls, explicit target selection, default and overridden server binding, and removal of remote stylesheet links. Existing tests exercise API routing, device categories, Cortex behavior, registration, workflow, and origin protection. JavaScript tests cover status updates, safe text rendering, and temporal navigation.')
p('The Python run emitted three warnings: two helper classes with constructors were not collected as test classes, and the modernization suite emitted its existing initialization warning. There were no test failures.')
h('Reproduction commands')
p('python -m pytest tests -q --basetemp=deliverables/sonarqube_fix/pytest-full')
p('node --test tests/test_dashboard_security.cjs tests/test_dashboard_temporal.cjs')
h('Required follow-up')
p('Run the configured SonarQube project again and reconcile the 19 bug/vulnerability findings and 6 related code smells against the new results. These are source-level remediations, not server-confirmed closures. No post-fix quality-gate status or zero-bug claim is made for the complete workbook.')
p('Review the 615 remaining in-scope code smells separately, especially complexity and exception handling. Third-party browser findings should be handled through vendor updates or a separate dependency review; backup findings remain historical. The current scan exclusions were not expanded to improve the totals.')
p('Before deployment, confirm the chosen LAN bind address where remote access is needed, exercise keyboard navigation and dashboard layout in the target browser, and verify MQTT, serial, and headset flows with the actual devices.')
h('Change and evidence records')
p('Eight application/configuration files changed: app.py, .env.example, core/devices.py, docs/generate_robotics_iot_diagrams.py, templates/dashboard.html, templates/cortex_classic.html, static/masterhub.css, and static/masterhub.js. Added tests/test_sonarqube_fixes.py and extended tests/test_dashboard_security.cjs.')
p('The sonarqube_fix deliverable folder retains original copies of the eight files, a unified patch, test outputs, and finding_disposition.csv with all 4,805 workbook rows classified. No Git metadata was present in the workspace.')
d.add_page_break()
h('Bug and vulnerability register')
p('All entries below are implemented in source and await SonarQube confirmation. Lines refer to the supplied workbook, because editing shifts source line numbers. Dashboard entries refer to templates/dashboard.html.')
records=[]
for r in rows:
 if not excluded(r[3]) and r[2] in ('BUG','VULNERABILITY'):
  f=r[3];rule=r[5]
  if f=='app.py':loc='app.py';fix='Default bind to localhost'
  elif f=='core/devices.py':loc='core/devices.py';fix='One normalized assignment'
  elif f.startswith('docs/'):loc='Robotics diagram generator';fix='Always use arc3 radius'
  elif rule=='Web:S5725':loc='Cortex Classic' if 'cortex' in f else 'Dashboard';fix='Remove remote font stylesheet'
  elif rule=='Web:S7930':loc='Dashboard';fix='Remove duplicate '+r[6].split('"')[1]
  else:loc='Dashboard';fix='Native keyboard button'
  short=rule.replace('Web:MouseEventWithoutKeyboardEquivalentCheck','Mouse keyboard equivalent')
  records.append([loc+'\nline '+r[4],short,fix])
table(['Source','Rule','Correction'],records,[1.6,1.9,3.4])

for node in list(d.styles.element.iter(qn('w:pBdr'))): node.getparent().remove(node)
for node in list(d.element.iter(qn('w:pBdr'))): node.getparent().remove(node)
d.save(out/'MasterHub_IoT_Bugfix_Report.docx')
print('Created report and full disposition register.')
