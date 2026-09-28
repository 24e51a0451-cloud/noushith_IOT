from pathlib import Path
import re
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

root = Path(__file__).parent
doc = Document()
sec = doc.sections[0]
sec.top_margin = sec.bottom_margin = Inches(.7)
sec.left_margin = sec.right_margin = Inches(.8)
sec.page_width = Inches(8.27)
sec.page_height = Inches(11.69)
for name in ['Normal', 'Title', 'Subtitle', 'Heading 1', 'Heading 2', 'Heading 3']:
    st = doc.styles[name]
    st.font.name = 'Calibri'
    st.font.color.rgb = RGBColor(0, 0, 0)
doc.styles['Normal'].font.size = Pt(11)
doc.styles['Normal'].paragraph_format.space_after = Pt(7)
doc.styles['Normal'].paragraph_format.line_spacing = 1.08
doc.styles['Title'].font.size = Pt(27)
doc.styles['Subtitle'].font.size = Pt(15)
for name, size in [('Heading 1', 17), ('Heading 2', 12)]:
    doc.styles[name].font.size = Pt(size)
    doc.styles[name].paragraph_format.space_before = Pt(12)
    doc.styles[name].paragraph_format.space_after = Pt(6)
    doc.styles[name].paragraph_format.keep_with_next = True

code = False
for line in (root/'report.md').read_text(encoding='utf-8').splitlines():
    if line.startswith('```'):
        code = not code
        continue
    if not line.strip():
        continue
    if code:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1
        r = p.add_run(line)
        r.font.name = 'Consolas'
        r.font.size = Pt(8.5)
    elif line.startswith('# '):
        doc.add_paragraph(line[2:], 'Title')
    elif line.startswith('## Research'):
        doc.add_paragraph(line[3:], 'Subtitle')
    elif line.startswith('## '):
        doc.add_paragraph(line[3:], 'Heading 1')
    elif line.startswith('### '):
        doc.add_paragraph(line[4:], 'Heading 2')
    else:
        p = doc.add_paragraph(line)
        if line.startswith('[S') or line.startswith('[R'):
            p.paragraph_format.keep_together = True
            for r in p.runs: r.font.size = Pt(10)
footer = sec.footer.paragraphs[0]
footer.alignment = 2
r = footer.add_run('MasterHub temporal framing   |   ')
r.font.size = Pt(9)
fld = OxmlElement('w:fldSimple')
fld.set(qn('w:instr'), 'PAGE')
footer._p.append(fld)
doc.core_properties.title = 'Temporal window framing in MasterHub'
doc.core_properties.subject = 'Manual and BCI command research and implementation guide'
doc.core_properties.author = 'MasterHub technical research'
for element in [doc._element, doc.styles.element]:
    for border in list(element.iter(qn('w:pBdr'))):
        border.getparent().remove(border)
out = root/'MasterHub_Temporal_Window_Research.docx'
doc.save(out)
print(out)
print('Words', len((root/'report.md').read_text(encoding='utf-8').split()))
