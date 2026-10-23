from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parent
doc = Document()
sec = doc.sections[0]
sec.page_height = Inches(11.69)
sec.page_width = Inches(8.27)
sec.top_margin = sec.bottom_margin = Inches(.7)
sec.left_margin = sec.right_margin = Inches(.75)
normal = doc.styles['Normal']
normal.font.name = 'Calibri'
normal.font.size = Pt(11)
normal.paragraph_format.space_after = Pt(7)
normal.paragraph_format.line_spacing = 1.08
for name, size in [('Title', 26), ('Subtitle', 12), ('Heading 1', 18), ('Heading 2', 12)]:
    s = doc.styles[name]
    s.font.name = 'Calibri'
    s.font.size = Pt(size)
    s.font.color.rgb = RGBColor(0,0,0)
    s.paragraph_format.space_after = Pt(9)
    s.paragraph_format.keep_with_next = True
code_style = doc.styles.add_style('Command', 1)
code_style.font.name = 'Consolas'
code_style.font.size = Pt(9)
code_style.paragraph_format.space_before = Pt(4)
code_style.paragraph_format.space_after = Pt(9)
code_style.paragraph_format.line_spacing = 1.0
code_style.paragraph_format.keep_together = True
foot = sec.footer.paragraphs[0]
foot.alignment = 2
foot.add_run('MasterHub Raspberry Pi 5 deployment  |  ').font.size = Pt(9)
field = OxmlElement('w:fldSimple'); field.set(qn('w:instr'), 'PAGE'); foot._p.append(field)
doc.core_properties.title = 'MasterHub Raspberry Pi 5 Deployment Guide'
doc.core_properties.subject = 'Step by step installation dashboard access and production service setup'
doc.core_properties.author = 'MasterHub'

def p(t, style=None): return doc.add_paragraph(t, style)
def h(t): doc.add_heading(t, 2)
def page(t): doc.add_page_break(); doc.add_heading(t, 1)
def code(t):
    x=p(t.strip(), 'Command')
    shade=OxmlElement('w:shd'); shade.set(qn('w:fill'),'F2F2F2'); x._p.get_or_add_pPr().append(shade)
def bullet(t): p(t, 'List Bullet')
def label(t):
    x=p(t); x.runs[0].bold=True

doc.add_heading('MasterHub Raspberry Pi 5 Deployment Guide', 0)
p('Step by step installation and production operation', 'Subtitle')
p('Prepared for the MasterHub team • 16 October 2026')
p('Run MasterHub continuously on Raspberry Pi 5 and open its dashboard from a laptop or phone. This guide covers installation, the existing PostgreSQL connection, encrypted dashboard access, automatic startup, HTTPS access for the team, device connections and operational checks.')
p('Start by keeping the existing PostgreSQL server and MQTT broker. Move the application first. An independent installation with PostgreSQL on the Pi requires a separate, verified database migration; the release ZIP does not contain your live database.')
h('Deployment status')
p('This is a procedure to execute on the Pi. It does not mean the Pi is already configured. Windows validation previously passed 250 backend tests and 17 frontend tests; ARM compatibility, HTTPS proxy changes and live hardware behavior still require validation on the actual Pi.')
h('Example names used throughout')
bullet('Pi hostname: masterhub. Pi administrator: noushith. Replace the administrator name if you choose another account.')
bullet('Application directory: /opt/masterhub. Background service user: masterhub. These are separate from the browser operator login.')
bullet('Example Pi address: 192.168.1.50. Replace it with the address assigned by your router.')
bullet('Initial dashboard: http://127.0.0.1:15000/dashboard through an SSH tunnel.')
bullet('Team dashboard: https://masterhub.yourcompany.com/dashboard after DNS, certificates and the HTTPS proxy are configured. This is a placeholder hostname.')
h('Before starting')
p('Use a Raspberry Pi 5 with a reliable power supply, cooling, Ethernet where possible, and reliable storage. SSD storage is preferable for sustained database and log writes. Keep a current database backup outside the Pi. Plan a maintenance window before switching real devices from the Windows controller.')
p('Follow steps 1 to 5 for the first working service, steps 6 and 7 for HTTPS and device cutover, and steps 8 and 9 for operations and troubleshooting.')

page('1 Prepare and connect to the Pi')
p('Use Raspberry Pi Imager to install Raspberry Pi OS Lite 64 bit. During imaging, set the hostname to masterhub, create your administrator account, select the timezone, and enable SSH. Prefer SSH key authentication. These steps do not require a desktop environment on the Pi.')
bullet('Connect Ethernet and power up the Pi. Reserve its IP address in the router so the address remains stable.')
bullet('The administrator account below is noushith. Use the account you actually created; there is no assumed default Pi password.')
label('Run on Windows in PowerShell')
code('ssh noushith@masterhub.local')
p('If the hostname does not resolve, find the Pi in the router device list and connect using its IP address.')
code('ssh noushith@192.168.1.50')
label('Run on the Pi after connecting')
code('hostname -I\nuname -m\npython3 --version')
p('Confirm a 64 bit system with aarch64 architecture. For this procedure use Python 3.12 or newer. If the OS provides an older Python, resolve the OS or interpreter choice before continuing; do not overwrite the system Python installation.')
h('Install the base packages')
code('sudo apt update\nsudo apt upgrade -y\nsudo apt install -y python3 python3-venv python3-pip unzip curl')
p('If the operating system requests a reboot, reboot and reconnect before continuing. Node.js is not required to serve the existing dashboard; it is only needed when running the frontend test suite.')
h('Check before proceeding')
bullet('SSH access works from your laptop.')
bullet('You know the Pi hostname, reserved IP address and administrator account.')
bullet('The Pi has network access to obtain dependencies and reach the intended database and MQTT broker.')
p('Reference: Raspberry Pi Getting started and remote access documentation, listed in step 10.')

page('2 Transfer and install MasterHub')
p('Use the compact release masterhub.zip from the project dist folder. Transfer source and static assets only. Windows virtual environments, Android SDKs, test caches and report previews do not belong in the Pi application installation.')
label('Run on Windows in PowerShell')
code('scp "D:\\GALATICX\\masterhub_iot_noushith\\dist\\masterhub.zip" `\n  noushith@masterhub.local:~/masterhub.zip')
p('The backtick at the end of the first line is PowerShell line continuation. Alternatively, enter the command on one line.')
label('Run on the Pi for a new installation')
code('sudo useradd --system --user-group \\\n  --home-dir /opt/masterhub \\\n  --shell /usr/sbin/nologin masterhub\nsudo mkdir -p /opt/masterhub\nsudo unzip ~/masterhub.zip -d /opt/masterhub\nsudo chown -R masterhub:masterhub /opt/masterhub')
p('If the service account or application directory already exists, inspect that installation instead of overwriting it. These commands assume a first installation. Keep the dedicated account without sudo privileges.')
h('Create the Linux virtual environment')
code('sudo -u masterhub python3 -m venv /opt/masterhub/.venv\nsudo -u masterhub /opt/masterhub/.venv/bin/python \\\n  -m pip install --upgrade pip\nsudo -u masterhub /opt/masterhub/.venv/bin/python \\\n  -m pip install -r /opt/masterhub/requirements-production.txt\nsudo -u masterhub /opt/masterhub/.venv/bin/python \\\n  -m pip check')
p('Expected dependency result: No broken requirements found. If a pinned dependency fails on ARM, stop and resolve that package error before enabling the service. Do not copy the Windows .venv to work around an installation failure.')
h('Create the private configuration file')
code('sudo -u masterhub cp /opt/masterhub/.env.example /opt/masterhub/.env\nsudo chmod 600 /opt/masterhub/.env\nsudo nano /opt/masterhub/.env')
p('The service user owns this file because the existing Cortex settings feature can save credentials. Never commit it or include it in a shared release ZIP. Populate it using the next page.')

page('3 Configure the database and operator login')
p('Replace every placeholder below. Keep the application listening on loopback for SSH tunnel access and the later HTTPS proxy. The operator password must be a unique random value of at least 32 characters.')
label('Contents to configure in the Pi .env file')
code('MASTERHUB_HOST=127.0.0.1\nMASTERHUB_PORT=5000\nMASTERHUB_OPERATOR_USER=operator\nMASTERHUB_OPERATOR_PASSWORD=REPLACE_WITH_LONG_RANDOM_PASSWORD\n\nPGHOST=YOUR_EXISTING_DATABASE_SERVER_ADDRESS\nPGPORT=YOUR_EXISTING_DATABASE_PORT\nPGDATABASE=masterhub\nPGUSER=masterhub_app\nPGPASSWORD=YOUR_EXISTING_DATABASE_PASSWORD\nDATABASE_URL=\nMASTERHUB_INSTANCE_ID=YOUR_EXISTING_INSTANCE_ID\n\nMQTT_HOST=YOUR_EXISTING_MQTT_BROKER\nMQTT_PORT=1883\nBCI_MOBILE_MQTT_HOST=YOUR_EXISTING_MQTT_BROKER\nBCI_MOBILE_MQTT_PORT=1883')
label('Generate a browser operator password on the Pi')
code('python3 -c "import secrets; print(secrets.token_urlsafe(32))"')
p('Save the generated password securely and paste it into MASTERHUB_OPERATOR_PASSWORD. The example MQTT port assumes the current non-TLS broker arrangement; retain the actual broker settings and configure its security before production exposure.')
h('Make the existing database reachable')
p('localhost on the Pi means the Pi, not your Windows computer. Set PGHOST to the existing database server address reachable from the Pi and preserve the actual port. The database administrator may need to adjust the server listener, PostgreSQL access rules and firewall to allow the Pi specifically. Use the server’s required TLS settings; never expose PostgreSQL publicly.')
p('A nonempty DATABASE_URL overrides the PG fields. Preserve the current MASTERHUB_INSTANCE_ID so historical events remain associated with this installation. If PostgreSQL stays on Windows, that computer must remain running.')
h('Validate before starting the app')
code('cd /opt/masterhub\nsudo -u masterhub .venv/bin/python tools/check_postgresql.py')
p('Expected result: connectivity, table initialization and a write/read/rollback check pass. If the check fails, fix connectivity, credentials or permissions before continuing. Database migration to the Pi is a separate backup, restore and verification operation; do not substitute an empty database for existing data.')

page('4 Start MasterHub and open the dashboard')
p('Before starting against real devices, ensure the old MasterHub controller is stopped or use isolated test endpoints. Two controllers can conflict even when the dashboard looks normal.')
label('Run on the Pi and leave the terminal open')
code('cd /opt/masterhub\nsudo -u masterhub .venv/bin/python serve.py')
p('Use serve.py for the production server. Do not deploy using Flask debug mode or its development server. The process requires a working PostgreSQL connection and configured operator credentials.')
h('Open an encrypted tunnel from Windows')
p('Open a second Windows PowerShell window. Keep the first Pi session running. Enter:')
code('ssh -N -L 15000:127.0.0.1:5000 noushith@masterhub.local')
p('No output is normal while the tunnel is running. Leave this PowerShell window open. If local port 15000 is occupied, choose another free local port and use the same number in the browser address.')
h('Open Chrome or Edge on your Windows laptop')
label('http://127.0.0.1:15000/dashboard')
p('Enter the operator username and password configured in the Pi .env file. The dashboard is served by the Pi; the laptop’s loopback address forwards through the encrypted SSH connection. No Python installation is needed in the browser.')
h('Check the service and data')
bullet('http://127.0.0.1:15000/healthz checks whether the web process responds.')
bullet('http://127.0.0.1:15000/api/health shows protected application health.')
bullet('Check Saved Activity in the dashboard to verify database access and persisted events.')
bullet('Confirm that connection indicators match the actual device state. An HTTP 200 response alone does not prove that hardware or storage is ready.')
h('What happens when you close a window')
p('Closing the tunnel disconnects the laptop browser. Closing the manual server session can stop the application. The next step replaces manual startup with a background service that survives SSH logout and starts at boot.')
p('A monitor does not need to remain attached to the Pi. Phone access without a tunnel is covered by the HTTPS setup in steps 6 and 7.')

page('5 Enable automatic startup with systemd')
p('Stop the manually started server using Ctrl+C in its Pi terminal. Keep only one MasterHub process running; device connections and control state are held in memory.')
label('Run on the Pi')
code('sudo nano /etc/systemd/system/masterhub.service')
label('Paste this service definition')
code('[Unit]\nDescription=MasterHub control service\nWants=network-online.target\nAfter=network-online.target\nStartLimitIntervalSec=0\n\n[Service]\nType=simple\nUser=masterhub\nGroup=masterhub\nWorkingDirectory=/opt/masterhub\nExecStart=/opt/masterhub/.venv/bin/python /opt/masterhub/serve.py\nEnvironment=PYTHONUNBUFFERED=1\nRestart=on-failure\nRestartSec=10\nTimeoutStopSec=30\nKillSignal=SIGTERM\nNoNewPrivileges=true\nUMask=0077\n\n[Install]\nWantedBy=multi-user.target')
p('The app loads its own .env file. This service retries a failed startup every 10 seconds, which allows recovery when the database is temporarily unavailable. It does not prove readiness or replace monitoring.')
code('sudo systemctl daemon-reload\nsudo systemctl enable --now masterhub\nsudo systemctl status masterhub --no-pager')
p('Expected result: Active: active (running). The existing SSH tunnel can now open the same dashboard even after the server SSH session is closed.')
h('Reboot check')
code('sudo reboot')
p('After reconnecting, run systemctl is-enabled masterhub, systemctl is-active masterhub, and curl http://127.0.0.1:5000/healthz. Reopen the laptop SSH tunnel after a Pi reboot.')

page('6 Configure HTTPS access for the team')
p('For browser access from laptops and phones without SSH tunnels, use Nginx with HTTPS. First obtain a real hostname, DNS pointing to the Pi’s reserved address, and a certificate trusted by the team devices. Replace masterhub.yourcompany.com and the certificate paths below. Do not publish the example hostname.')
h('Required application change before using the proxy')
p('The current origin protection must see the original HTTPS scheme. Otherwise the dashboard can load while button actions fail with a cross-origin error. Add these arguments to the existing Waitress serve call in serve.py; this guide does not apply the change automatically:')
code('trusted_proxy="127.0.0.1",\ntrusted_proxy_headers={"x-forwarded-proto"},')
p('Keep MASTERHUB_HOST=127.0.0.1. Trust only the local proxy for this header and preserve operator authentication. Validate this change with login and a safe state-changing request.')
label('Run on the Pi')
code('sudo apt install -y nginx\nsudo nano /etc/nginx/sites-available/masterhub')
label('Paste after providing the certificate files')
code('server {\n    listen 443 ssl;\n    server_name masterhub.yourcompany.com;\n    ssl_certificate /etc/nginx/tls/masterhub-fullchain.pem;\n    ssl_certificate_key /etc/nginx/tls/masterhub-key.pem;\n    ssl_protocols TLSv1.2 TLSv1.3;\n    client_max_body_size 1m;\n\n    location / {\n        proxy_pass http://127.0.0.1:5000;\n        proxy_http_version 1.1;\n        proxy_set_header Host $http_host;\n        proxy_set_header X-Forwarded-Proto $scheme;\n        proxy_set_header Connection "";\n        proxy_buffering off;\n        proxy_read_timeout 3600s;\n    }\n}')
p('Keep the private key restricted to the administrator and Nginx service requirements. Disabling buffering supports the live event stream. The long read timeout allows quiet streaming connections; validate stream reconnects in your environment.')

page('7 Enable team access and reconnect devices')
label('Enable the HTTPS site on the Pi')
code('sudo ln -s /etc/nginx/sites-available/masterhub \\\n  /etc/nginx/sites-enabled/masterhub\nsudo nginx -t\nsudo systemctl restart masterhub\nsudo systemctl enable nginx\nsudo systemctl reload nginx')
p('If the site link already exists, do not recreate it. If nginx -t fails, fix the hostname, syntax or certificate paths before reloading. Allow HTTPS from the intended network and retain SSH access. Keep application port 5000 and PostgreSQL closed to ordinary browser clients.')
h('Where the team opens the dashboard')
label('https://masterhub.yourcompany.com/dashboard')
p('Replace this placeholder with the real configured hostname. Open it in a laptop or phone browser on the permitted network and enter the operator credentials. For access outside the office, use the organisation’s VPN unless a public deployment is deliberately designed and reviewed.')
h('Reconnect each endpoint')
bullet('Windows PC: keep the PC agent running in its interactive Windows session. Select that PC as the target. Local Host now means the Pi and cannot run Windows Notepad or Windows desktop automation.')
bullet('Android: point the companion app at the intended broker and registration endpoint. Native control clients need operator HTTP Basic credentials; Android account endpoints retain their existing bearer-token contract.')
bullet('IoT hardware: retain exact MQTT topics and broker settings. Saved .runtime/hardware.json can override environment defaults. Review device-specific MQTT authentication and permissions before production exposure.')
bullet('USB or UART: reconnect the device and select its Linux serial port. Use a persistent device identity where possible rather than assuming COM port names carry over.')
code('sudo usermod -aG dialout masterhub\nsudo systemctl restart masterhub')
p('Only add serial access if required. EMOTIV Cortex must be configured separately: localhost:6868 now refers to the Pi. Confirm headset, OS and Cortex compatibility, or use the supported remote Cortex connection. Do not simply disable TLS verification to connect to a remote host.')
h('Cut over without losing runtime data')
p('Stop actuators safely, then stop the old Windows controller before making the Pi the live controller. Preserve MASTERHUB_INSTANCE_ID, data/event-outbox and relevant .runtime/hardware.json settings. Transfer pending events only while the old writer is stopped. Keep backups; the release ZIP excludes live data.')

page('8 Operate and verify the production service')
label('Service commands on the Pi')
code('sudo systemctl status masterhub --no-pager\nsudo systemctl restart masterhub\nsudo systemctl stop masterhub\nsudo systemctl start masterhub\nsudo journalctl -u masterhub -n 100 --no-pager\nsudo journalctl -u masterhub -f')
p('Press Ctrl+C to stop following logs; that does not stop the background service. Restart after changing .env. Run daemon-reload as well if the service definition changes.')
h('Production acceptance sequence')
for t in [
'1. Confirm dashboard login, CSS loading and accurate device status through the final access method.',
'2. Check Saved Activity and verify that a new event is persisted in PostgreSQL.',
'3. Issue a safe command to the intended device and confirm its acknowledgement and actual behavior.',
'4. Restart MasterHub and verify that stale actuator commands are not replayed.',
'5. Reboot the Pi and confirm automatic startup, dashboard access and device reconnection.',
'6. Test database and MQTT outages in a controlled window. Confirm visible failures, queued event handling and recovery.',
'7. Run a sustained workload representative of the expected devices and browser sessions. Watch CPU, memory, temperature, disk use and command latency.',
'8. Restore a database backup into an isolated test database and verify the recovered records. A backup file existing is not proof that restoration works.'
]: p(t)
h('Backups and updates')
p('Schedule database backups on the database host and copy them to separate storage. Preserve private configuration securely, the event outbox, hardware settings and installation ID. Never package credentials into an ordinary release.')
p('For each update: pass tests, keep the previous release, stop actuators, back up data, stop the service, install the new verified release, preserve configuration and pending events, restart, and perform the acceptance checks. Avoid overwriting a running installation.')
p('If an update fails, stop the new process and restore the previous application release and compatible configuration. Database schema changes require their own rollback plan. Do not roll back by deleting queued events or replacing the active database with an older copy without reconciling new data.')

page('9 Troubleshoot common problems')
h('SSH cannot find the Pi')
p('Use its actual IP address instead of masterhub.local. Confirm the Pi is powered on, SSH was enabled, and both machines can reach the same network. Guest Wi-Fi may isolate devices.')
h('The dashboard does not open through the tunnel')
p('Keep the tunnel PowerShell window open. Confirm the local port is free, the service is active, and the Pi can answer curl http://127.0.0.1:5000/healthz. The browser uses port 15000 on the laptop, while the Pi listener uses port 5000.')
h('The service repeatedly restarts')
code('sudo journalctl -u masterhub -n 100 --no-pager')
p('Look for missing credentials, an operator password shorter than 32 characters, unreachable PostgreSQL, dependency import failures, or port 5000 already in use. Stop any manual serve.py process before starting the service.')
h('Database connection fails')
p('Check the real server address and port, database name, role permissions, password and TLS configuration. localhost cannot reach a database on another computer. Confirm that DATABASE_URL is not unexpectedly overriding the PG fields and that the server allows the Pi through its firewall and PostgreSQL access rules.')
h('HTTPS loads but buttons fail')
p('Confirm that the trusted-proxy arguments were added to Waitress and Nginx forwards the original Host and X-Forwarded-Proto. Validate the proxy configuration, restart MasterHub and test a safe mutation. Do not remove cross-origin protection to hide the error.')
h('Certificate warning in the browser')
p('Use the hostname covered by the certificate. Check the full certificate chain, certificate expiry, DNS and client trust. An internal certificate authority must be trusted on every intended client. Do not instruct operators to bypass certificate warnings.')
h('Devices appear offline or desktop commands fail')
p('Verify the broker address, topics, credentials and saved hardware settings. Ensure the PC or phone agent is running and selected. Windows automation needs the Windows agent and an interactive session; the Pi Local Host target does not provide it.')
h('Live activity stops updating')
p('Check the event stream connection, proxy buffering and timeouts, and browser reconnection behavior. The bounded Waitress thread pool is intended for a small operator group. Measure concurrent streams before expanding the deployment.')

page('10 References and handover checklist')
p('These sources support the platform setup and deployment concepts. Project-specific paths, credentials and device mappings must be checked against the actual installation.')
for title,url in [
('Raspberry Pi setup and imaging','https://www.raspberrypi.com/documentation/computers/getting-started.html'),
('Raspberry Pi hardware and power guidance','https://www.raspberrypi.com/documentation/computers/raspberry-pi.html'),
('Flask production deployment','https://flask.palletsprojects.com/en/stable/deploying/'),
('Flask Waitress deployment','https://flask.palletsprojects.com/en/stable/deploying/waitress/'),
('Nginx HTTPS configuration','https://nginx.org/en/docs/http/configuring_https_servers.html'),
('Nginx proxy configuration','https://nginx.org/en/docs/http/ngx_http_proxy_module.html'),
('PostgreSQL Debian packages and ARM64 support','https://www.postgresql.org/download/linux/debian/'),
('EMOTIV Cortex platform requirements','https://emotiv.gitbook.io/system-requirements/cortex-api'),
('systemd service reference','https://www.freedesktop.org/software/systemd/man/latest/systemd.service.html')]:
    label(title); p(url)
h('Record before handing the service to the team')
bullet('Pi hostname and reserved address; administrator and service owner.')
bullet('Final dashboard URL; certificate issuer and renewal owner.')
bullet('PostgreSQL host, port and backup owner; restore test date.')
bullet('MQTT broker and device topic owners; Cortex host and headset configuration.')
bullet('Release version, deployment date, acceptance results and rollback location.')
bullet('Credential storage location and recovery procedure, without writing passwords into this guide.')
p('Keep the Pi service account unprivileged, restrict network access, and monitor storage and connection failures. A single Pi remains a single point of failure; define the expected recovery time and safe device behavior if it loses power.')

out = ROOT / 'MasterHub_Raspberry_Pi_5_Deployment_Guide.docx'
for element in (doc.styles.element, doc.element):
    for border in list(element.iter(qn('w:pBdr'))):
        border.getparent().remove(border)
doc.save(out)
print(out)

